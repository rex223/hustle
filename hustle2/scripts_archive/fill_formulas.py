import openpyxl

file_path = 'RENTAL DATASHEET 1ST AUG.xlsm'
wb = openpyxl.load_workbook(file_path, keep_vba=True)

ignore_sheets = ['RENT DATA', 'PAYMENT SHEET', 'EXPENDITURE', 'PURCHASE', 'Sheet17']

for sheet_name in wb.sheetnames:
    if sheet_name in ignore_sheets:
        continue
        
    ws = wb[sheet_name]
    base_name = sheet_name.upper()
    
    # Fill formulas in rows 2 to 500
    for r in range(2, 501):
        ws.cell(row=r, column=4).value = base_name
        ws.cell(row=r, column=5).value = f'=SUMIF(\'RENT DATA\'!A:A, "{base_name}*", \'RENT DATA\'!C:C)'
        ws.cell(row=r, column=7).value = f'=VLOOKUP("{base_name}*", \'RENT DATA\'!A:E, 5, FALSE)'
        ws.cell(row=r, column=8).value = f'=E{r}-F{r}'

wb.save(file_path)
print("Formulas filled down to row 500 for all individual sheets.")
