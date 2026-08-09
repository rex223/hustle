import pandas as pd
import openpyxl
from openpyxl.styles import Font
import re

file_path = 'RENTAL DATASHEET 30 JULY.xlsm'
wb = openpyxl.load_workbook(file_path, keep_vba=True)

rent_ws = wb['RENT DATA']

# Extract all base names from Column A of RENT DATA
base_names = set()
for row in range(2, rent_ws.max_row + 1):
    val = rent_ws.cell(row=row, column=1).value
    if val:
        val = str(val).strip()
        if val.upper() == "GRAND TOTAL" or val.upper().startswith("TOTAL"):
            continue
        
        # Strip trailing numbers and spaces
        while re.search(r'[\d\s]$', val):
            val = val[:-1]
        
        if val:
            base_names.add(val.upper())

ignore_sheets = ['RENT DATA', 'PAYMENT SHEET', 'EXPENDITURE', 'PURCHASE', 'Sheet17']
bold_font = Font(bold=True)
headers = ['DATE', 'TIME', 'FOR THE MONTH', 'NAME', 'AMOUNT', 'PAID AMOUNT', 'GROUP', 'DUE']

# 1. Update existing sheets and create new ones
for base_name in base_names:
    if base_name in ignore_sheets:
        continue
        
    wsExists = False
    for sheet in wb.sheetnames:
        if sheet.upper() == base_name:
            wsExists = True
            ws = wb[sheet]
            break
            
    if not wsExists:
        ws = wb.create_sheet(title=base_name)
        
    # Apply Headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = bold_font
        
    # Inject formulas in row 2
    ws.cell(row=2, column=4).value = base_name
    ws.cell(row=2, column=5).value = f'=SUMIF(\'RENT DATA\'!A:A, "{base_name}*", \'RENT DATA\'!C:C)'
    # PAID AMOUNT is column 6 (left empty for manual entry)
    # GROUP is column 7
    ws.cell(row=2, column=7).value = f'=VLOOKUP("{base_name}*", \'RENT DATA\'!A:E, 5, FALSE)'
    # DUE is column 8
    ws.cell(row=2, column=8).value = '=E2-F2'
    
    column_widths = {'A': 15, 'B': 12, 'C': 20, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

wb.save(file_path)
print("Successfully synced all names to individual sheets!")
