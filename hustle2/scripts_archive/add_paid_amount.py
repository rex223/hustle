import openpyxl
from openpyxl.styles import Font

file_path = 'RENTAL DATASHEET 1ST AUG.xlsm'
wb = openpyxl.load_workbook(file_path, keep_vba=True)

ignore_sheets = ['RENT DATA', 'PAYMENT SHEET', 'EXPENDITURE', 'PURCHASE', 'Sheet17']
bold_font = Font(bold=True)

for sheet_name in wb.sheetnames:
    if sheet_name in ignore_sheets:
        continue
        
    ws = wb[sheet_name]
    
    # Check current headers to avoid double processing
    headers = [ws.cell(row=1, column=i).value for i in range(1, 10)]
    if 'PAID AMOUNT' not in headers:
        # Insert column at F (index 6)
        ws.insert_cols(6)
        
        # Set Header
        f1 = ws.cell(row=1, column=6)
        f1.value = 'PAID AMOUNT'
        f1.font = bold_font
        ws.column_dimensions['F'].width = 15
        
        # Original G (DUE) was moved to H, let's just make sure H1 is DUE
        h1 = ws.cell(row=1, column=8)
        h1.value = 'DUE'
        h1.font = bold_font
        
        # Add formula to DUE column at row 2
        # E is AMOUNT, F is PAID AMOUNT
        ws.cell(row=2, column=8).value = '=E2-F2'

wb.save(file_path)
print("Updated all individual sheets with PAID AMOUNT and DUE formula.")
