import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

# --- Config ---
EXCEL_PATH = Path(__file__).parent.parent / "sheets" / "forms_v5.xlsm"
LOCAL_CSV = Path(__file__).parent.parent / "sheets" / "local_transport.csv"

def main():
    if not EXCEL_PATH.exists() or not LOCAL_CSV.exists():
        print("Missing Excel or Local CSV file.")
        return
        
    print("Loading Local Transport Data...")
    local_df = pd.read_csv(LOCAL_CSV, dtype=str).fillna("")
    if local_df.empty:
        print("No local data to inject.")
        return
        
    local_dict = {}
    for _, row in local_df.iterrows():
        local_dict[row['CODE']] = {
            'auto': row['Auto Stand'],
            'bazaar': row['Main Bazaar']
        }
        
    print(f"Injecting {len(local_dict)} rows into Excel...")
    
    import warnings
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    
    wb = load_workbook(EXCEL_PATH, keep_vba=True)
    ws = wb["Database"]
    
    # Find columns dynamically just in case
    headers = {cell.value: i for i, cell in enumerate(ws[1]) if cell.value}
    code_col = headers.get('CODE')
    auto_col = headers.get('Auto Stand')
    bazaar_col = headers.get('Main Bazaar')
    
    if None in (code_col, auto_col, bazaar_col):
        print("Error: Could not find required columns in Excel sheet.")
        return
        
    injected = 0
    for row in ws.iter_rows(min_row=2):
        code = str(row[code_col].value).strip() if row[code_col].value else ""
        if code in local_dict:
            row[auto_col].value = local_dict[code]['auto']
            row[bazaar_col].value = local_dict[code]['bazaar']
            injected += 1
            
    if injected > 0:
        wb.save(EXCEL_PATH)
        print(f"Successfully injected {injected} rows and saved {EXCEL_PATH.name}!")
    else:
        print("No matching CODEs found to inject.")

if __name__ == "__main__":
    main()
