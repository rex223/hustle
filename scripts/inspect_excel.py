import openpyxl

def inspect_excel(filepath):
    wb = openpyxl.load_workbook(filepath, keep_vba=True)
    for sheet_name in wb.sheetnames:
        print(f"Sheet: {sheet_name}")
        ws = wb[sheet_name]
        
        # Assuming headers are in row 4 for Forms and row 1 for Database
        for row_idx in range(1, 6):
            row_vals = [str(cell.value) if cell.value else '' for cell in ws[row_idx]]
            if any(row_vals):
                print(f"Row {row_idx}: {row_vals}")

inspect_excel("forms.xlsm")
