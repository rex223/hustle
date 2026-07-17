import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

# --- Config ---
EXCEL_PATH = Path(__file__).parent.parent / "sheets" / "forms_v4.xlsm"
TRAIN_CSV = Path(__file__).parent.parent / "sheets" / "train_data.csv"

def main():
    if not EXCEL_PATH.exists() or not TRAIN_CSV.exists():
        print("Missing Excel or Train CSV file.")
        return
        
    print("Loading Train Data...")
    train_df = pd.read_csv(TRAIN_CSV, dtype=str).fillna("")
    if train_df.empty:
        print("No train data to inject.")
        return
        
    train_dict = {}
    for _, row in train_df.iterrows():
        train_dict[row['CODE']] = {
            'station': row['Nearest Railway Station'],
            'train': row['Train from Sealdah']
        }
        
    print(f"Injecting {len(train_dict)} rows into Excel...")
    
    import warnings
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    
    wb = load_workbook(EXCEL_PATH, keep_vba=True)
    ws = wb["Database"]
    
    # Find columns dynamically just in case
    headers = {cell.value: i for i, cell in enumerate(ws[1]) if cell.value}
    code_col = headers.get('CODE')
    station_col = headers.get('Nearest Railway Station')
    train_col = headers.get('Train from Sealdah')
    
    if None in (code_col, station_col, train_col):
        print("Error: Could not find required columns in Excel sheet.")
        return
        
    injected = 0
    for row in ws.iter_rows(min_row=2):
        code = str(row[code_col].value).strip() if row[code_col].value else ""
        if code in train_dict:
            row[station_col].value = train_dict[code]['station']
            row[train_col].value = train_dict[code]['train']
            injected += 1
            
    if injected > 0:
        wb.save(EXCEL_PATH)
        print(f"Successfully injected {injected} rows and saved {EXCEL_PATH.name}!")
    else:
        print("No matching CODEs found to inject.")

if __name__ == "__main__":
    main()
